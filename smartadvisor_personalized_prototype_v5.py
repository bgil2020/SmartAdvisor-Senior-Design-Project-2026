"""SmartAdvisor personalized Course Compatibility Engine prototype.

Place this script in the same folder as:
  - Synthetic Profiles.xlsx
  - Your latest Final_FAU_CS_CE_EE_Undergraduate_Courses_Revised*.xlsx
  - Course Combinations.xlsx

Requires: openpyxl
Install if needed: py -m pip install openpyxl

This prototype uses synthetic student profiles and course files for testing only.
It provides planning guidance, not an official advising or registration decision.
"""

from __future__ import annotations

import re
from pathlib import Path
from itertools import combinations
from difflib import SequenceMatcher
from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent
PROFILE_FILE = BASE_DIR / "Synthetic Profiles.xlsx"
CATALOG_CANDIDATES = sorted(
    BASE_DIR.glob("Final_FAU_CS_CE_EE_Undergraduate_Courses_Revised*.xlsx"),
    key=lambda p: p.stat().st_mtime,
    reverse=True,
)
CATALOG_FILE = CATALOG_CANDIDATES[0] if CATALOG_CANDIDATES else BASE_DIR / "Final_FAU_CS_CE_EE_Undergraduate_Courses_Revised.xlsx"
COMBINATIONS_FILE = BASE_DIR / "Course Combinations.xlsx"


def clean(value) -> str:
    return "" if value is None else str(value).strip()


def normalize_course_code(value: str) -> str:
    """Convert 'EEL 3111' and 'eel3111' to 'EEL3111'."""
    return re.sub(r"[^A-Z0-9]", "", clean(value).upper())


def display_course_code(code: str) -> str:
    m = re.match(r"^([A-Z]+)(\d+[A-Z]?)$", normalize_course_code(code))
    return f"{m.group(1)} {m.group(2)}" if m else clean(code).upper()


def normalize_text(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", clean(value).lower())


def parse_credit(value) -> float:
    text = clean(value)
    if not text:
        return 0.0
    nums = re.findall(r"\d+(?:\.\d+)?", text)
    if not nums:
        return 0.0
    # Catalog sometimes represents a range. Use the highest listed value.
    return max(float(x) for x in nums)


def read_sheet_rows(path: Path) -> list[dict]:
    if not path.exists():
        raise FileNotFoundError(f"Required file not found: {path.name}")
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    headers = [clean(x) for x in next(rows)]
    result = []
    for row in rows:
        if not any(v is not None and clean(v) for v in row):
            continue
        result.append({headers[i]: row[i] if i < len(row) else None for i in range(len(headers))})
    wb.close()
    return result


def load_profiles() -> dict[str, dict]:
    profiles = {}
    for row in read_sheet_rows(PROFILE_FILE):
        sid = clean(row.get("Student ID")).upper()
        if sid:
            profiles[sid] = {
                "student_id": sid,
                "major": clean(row.get("Major")),
                "name": clean(row.get("First and Last Name")),
                "standing": clean(row.get("Undergraduate Standing /Level")),
                "completed_courses_raw": clean(row.get("Completed Courses")),
                "workdays": clean(row.get("Workdays")),
                "work_schedule": clean(row.get("Work Schedule")),
            }
    return profiles


def load_catalog() -> dict[str, dict]:
    """Collapse multiple section rows into one course record while retaining sections."""
    catalog: dict[str, dict] = {}
    for row in read_sheet_rows(CATALOG_FILE):
        raw_code = clean(row.get("COURSE"))
        code = normalize_course_code(raw_code)
        if not code:
            continue
        credits = parse_credit(row.get("CREDIT HOURS LOW:HIGH"))
        record = catalog.setdefault(code, {
            "code": code,
            "display_code": raw_code or display_course_code(code),
            "title": clean(row.get("TITLE")),
            "credits": 0.0,
            "sections": [],
        })
        # Discussion/lab rows can be 0 credits, so keep the maximum course credit value.
        record["credits"] = max(record["credits"], credits)
        if not record["title"] and clean(row.get("TITLE")):
            record["title"] = clean(row.get("TITLE"))
        record["sections"].append({
            "section": clean(row.get("SECT")),
            "crn": clean(row.get("CRN")),
            "status": clean(row.get("ENRL STAT")),
            "campus": clean(row.get("CAMPUS")),
            "days": clean(row.get("DAYS")),
            "begin": clean(row.get("BEG TIME")),
            "end": clean(row.get("END TIME")),
            "remaining": row.get("ENRL REMN"),
        })
    return catalog


def resolve_course(raw_id: str, course_name: str, catalog: dict[str, dict]) -> str | None:
    """Resolve minor inconsistencies in the manually prepared combinations sheet."""
    code = normalize_course_code(raw_id)
    if code in catalog:
        return code

    # Handle a missing suffix such as COP3275 vs COP3275C when only one match exists.
    prefix_matches = [c for c in catalog if c.startswith(code) or code.startswith(c)] if code else []
    if len(prefix_matches) == 1:
        return prefix_matches[0]

    # If an ID cell accidentally contains a course name, try the Course Name column.
    target = normalize_text(course_name or raw_id)
    if not target:
        return None
    exact_title = [c for c, info in catalog.items() if normalize_text(info["title"]) == target]
    if len(exact_title) == 1:
        return exact_title[0]

    # Conservative fuzzy title fallback.
    scored = []
    for c, info in catalog.items():
        title = normalize_text(info["title"])
        if title:
            scored.append((SequenceMatcher(None, target, title).ratio(), c))
    scored.sort(reverse=True)
    if scored and scored[0][0] >= 0.88 and (len(scored) == 1 or scored[0][0] - scored[1][0] >= 0.05):
        return scored[0][1]
    return None


def load_course_combinations(catalog: dict[str, dict]) -> tuple[dict[frozenset, dict], list[str]]:
    pair_data = {}
    warnings = []
    for row in read_sheet_rows(COMBINATIONS_FILE):
        a = resolve_course(clean(row.get("Course Id")), clean(row.get("Course Name")), catalog)
        b = resolve_course(clean(row.get("Course B ID")), clean(row.get("Course B Name")), catalog)
        if not a or not b:
            warnings.append(
                f"Could not fully match combination: {clean(row.get('Course Name'))} / {clean(row.get('Course B Name'))}"
            )
            continue
        try:
            rating = float(row.get("Rating Out Of 10"))
        except (TypeError, ValueError):
            rating = None
        pair_data[frozenset((a, b))] = {
            "course_a": a,
            "course_b": b,
            "reasoning": clean(row.get("Reasoning")),
            "baseline_rating": rating,
        }
    return pair_data, warnings


def extract_completed_codes(text: str, catalog: dict[str, dict]) -> list[str]:
    """Extract recognizable course codes even when commas are missing in the synthetic file."""
    compact = clean(text).upper()
    candidates = re.findall(r"[A-Z]{2,4}\s*\d{4}[A-Z]?", compact)
    found = []
    for candidate in candidates:
        code = normalize_course_code(candidate)
        if code in catalog and code not in found:
            found.append(code)
    return found


def prompt_int(prompt: str, minimum: int, maximum: int, allow_blank=False, default=None) -> int | None:
    while True:
        raw = input(prompt).strip()
        if not raw and allow_blank:
            return default
        try:
            value = int(raw)
            if minimum <= value <= maximum:
                return value
        except ValueError:
            pass
        print(f"Please enter a number from {minimum} to {maximum}.")


def prompt_float(prompt: str, minimum=0.0, maximum=100.0, allow_blank=False, default=None) -> float | None:
    while True:
        raw = input(prompt).strip()
        if not raw and allow_blank:
            return default
        try:
            value = float(raw)
            if minimum <= value <= maximum:
                return value
        except ValueError:
            pass
        print(f"Please enter a number from {minimum:g} to {maximum:g}.")


def prompt_yes_no(prompt: str) -> bool:
    while True:
        raw = input(prompt + " (Y/N): ").strip().lower()
        if raw in {"y", "yes"}:
            return True
        if raw in {"n", "no"}:
            return False
        print("Please enter Y or N.")


def select_profile(profiles: dict[str, dict]) -> dict:
    print("\nAvailable synthetic profiles:")
    for sid, p in profiles.items():
        print(f"  {sid}: {p['name']} — {p['major']} ({p['standing']})")
    while True:
        sid = input("\nEnter synthetic Student ID: ").strip().upper()
        if sid in profiles:
            return profiles[sid]
        print("Student ID not found. Please choose one shown above.")


def select_recent_courses(catalog: dict[str, dict], completed: list[str]) -> list[str]:
    """Ask which courses were taken together in the student's most recent semester."""
    print("\nRECENT SEMESTER COURSES")
    if completed:
        print("Recognized completed courses in this synthetic profile:")
        print("  " + ", ".join(display_course_code(c) for c in completed))
    print("Enter the courses taken in the most recent semester, separated by commas.")
    print("Press Enter if this information is unknown.")

    while True:
        raw = input("Recent-semester courses: ").strip()
        if not raw:
            return []
        entries = [x.strip() for x in re.split(r"[,;]", raw) if x.strip()]
        selected, missing = [], []
        for entry in entries:
            code = normalize_course_code(entry)
            if code in catalog:
                if code not in selected:
                    selected.append(code)
            else:
                matches = [c for c in catalog if c.startswith(code)]
                if len(matches) == 1:
                    selected.append(matches[0])
                else:
                    missing.append(entry)
        if missing:
            print("Not found in the current catalog: " + ", ".join(missing))
            print("Please try again, or press Enter if the recent schedule is unknown.")
            continue
        return selected


def select_courses(catalog: dict[str, dict]) -> list[str]:
    print("\nEnter the courses you want SmartAdvisor to evaluate.")
    print("You may enter them together (example: EEL3111, EEL3502, STA4821).")
    while True:
        raw = input("Proposed courses: ").strip()
        entries = [x.strip() for x in re.split(r"[,;]", raw) if x.strip()]
        selected, missing = [], []
        for entry in entries:
            code = normalize_course_code(entry)
            if code in catalog:
                if code not in selected:
                    selected.append(code)
            else:
                # Accept a unique prefix if the user omitted a catalog suffix such as C.
                matches = [c for c in catalog if c.startswith(code)]
                if len(matches) == 1:
                    selected.append(matches[0])
                else:
                    missing.append(entry)
        if missing:
            print("Not found in the current catalog: " + ", ".join(missing))
            print("Please re-enter the proposed schedule using courses in the catalog.")
            continue
        if not selected:
            print("Please enter at least one course.")
            continue
        return selected


def baseline_label(rating: float | None) -> str:
    """The manual 1-10 rating is used as a baseline input, not a scientific risk score."""
    if rating is None:
        return "Not rated"
    if rating >= 8:
        return "Elevated consideration"
    if rating >= 6:
        return "Moderate consideration"
    return "Lower consideration"


def personalized_pair_guidance(pair: dict, difficulty: int, satisfied: str,
                               current_work_hours: float, proposed_credits: float,
                               preferred_max: float | None,
                               recent_courses: list[str] | None = None,
                               proposed_courses: list[str] | None = None) -> tuple[str, list[str]]:
    """Translate baseline pair information + student context into cautious guidance.

    This is intentionally rule-based and categorical. It does not claim that the
    spreadsheet's 1-10 baseline rating is an empirically validated probability.
    """
    factors = []
    concern = 0

    rating = pair.get("baseline_rating")
    if rating is not None:
        if rating >= 8:
            concern += 2
        elif rating >= 6:
            concern += 1

    if difficulty >= 5:
        concern += 2
        factors.append("The student reported that a recent workload felt very difficult.")
    elif difficulty == 4:
        concern += 1
        factors.append("The student reported that a recent workload felt difficult.")
    elif difficulty == 1:
        concern -= 1
        factors.append("The student reported that the recent semester felt very manageable.")
    elif difficulty == 2:
        concern -= 1
        factors.append("The student reported that the recent semester felt manageable.")

    if satisfied == "no":
        concern += 2
        factors.append("The student was not satisfied with performance during that recent workload.")
    elif satisfied == "mostly":
        concern += 1
        factors.append("The student was only partly satisfied with performance during that recent workload.")
    else:
        factors.append("The student was satisfied with performance during that recent workload.")

    if current_work_hours >= 20:
        concern += 2
        factors.append(f"The student expects to work about {current_work_hours:g} hours per week.")
    elif current_work_hours >= 10:
        concern += 1
        factors.append(f"The student expects to work about {current_work_hours:g} hours per week.")

    if recent_courses and proposed_courses:
        repeated = sorted(set(recent_courses) & set(proposed_courses))
        if repeated:
            factors.append(
                "The student has recent experience with " +
                ", ".join(display_course_code(c) for c in repeated) + "."
            )

    if preferred_max is not None and proposed_credits > preferred_max:
        concern += 2
        factors.append(
            f"The proposed {proposed_credits:g}-credit schedule is above the student's preferred maximum of {preferred_max:g} credits."
        )

    if concern >= 6:
        result = "Potential compatibility concern"
    elif concern >= 3:
        result = "Review recommended"
    else:
        result = "Potentially manageable with consideration"
    return result, factors


def overall_semester_guidance(proposed_credits: float, preferred_max: float | None,
                              difficulty: int, satisfied: str, work_hours: float,
                              matched_pairs: int) -> str:
    concern = 0
    if preferred_max is not None:
        if proposed_credits > preferred_max:
            concern += 3
        elif proposed_credits == preferred_max:
            concern += 1
    if difficulty >= 4:
        concern += 1
    if satisfied in {"mostly", "no"}:
        concern += 1
    if work_hours >= 20:
        concern += 2
    elif work_hours >= 10:
        concern += 1
    if matched_pairs >= 2:
        concern += 1

    if concern >= 6:
        return "Heavy relative to the student's stated circumstances"
    if concern >= 3:
        return "Moderate-to-heavy relative to the student's stated circumstances"
    return "Potentially manageable based on the information provided"



def build_personalized_recommendation(proposed_credits: float, recent_credits: float | None,
                                      preferred_max: float | None, difficulty: int,
                                      satisfied: str, work_hours: float,
                                      matched_pairs: list[dict], recent_courses: list[str],
                                      proposed_courses: list[str]) -> list[str]:
    """Build a recommendation from the student's responses and stored course data."""
    recommendations = []
    concerns = []
    strengths = []

    # Compare the proposed load with the student's own recent experience.
    if recent_credits is not None:
        difference = proposed_credits - recent_credits
        if difference >= 3:
            concerns.append(
                f"The proposed schedule is {difference:g} credits higher than the student's recent {recent_credits:g}-credit semester."
            )
        elif difference > 0:
            concerns.append(
                f"The proposed schedule is slightly higher than the student's recent {recent_credits:g}-credit semester."
            )
        elif difference <= -3:
            strengths.append(
                f"The proposed schedule is {abs(difference):g} credits lower than the student's recent {recent_credits:g}-credit semester."
            )
        else:
            strengths.append("The proposed credit load is similar to the student's recent semester.")

    # Use the student's own description of the recent semester.
    if difficulty >= 4:
        concerns.append("The student reported that the recent semester felt difficult.")
    elif difficulty == 1:
        strengths.append("The student reported that the recent semester felt very manageable.")
    elif difficulty == 2:
        strengths.append("The student reported that the recent semester felt manageable.")

    if satisfied == "no":
        concerns.append("The student was not satisfied with academic performance during that recent workload.")
    elif satisfied == "mostly":
        concerns.append("The student was only partly satisfied with academic performance during that recent workload.")
    elif satisfied == "yes":
        strengths.append("The student was satisfied with academic performance during that recent workload.")

    # Current commitments.
    if work_hours >= 20:
        concerns.append(f"The student expects to work about {work_hours:g} hours per week during the proposed semester.")
    elif work_hours >= 10:
        recommendations.append(
            f"Plan study time around the expected {work_hours:g} work hours per week, especially during exam and project weeks."
        )

    # Student's own preferred maximum should matter directly.
    if preferred_max is not None:
        if proposed_credits > preferred_max:
            recommendations.append(
                f"Consider reducing the proposed load to {preferred_max:g} credits or below because it exceeds the student's stated preferred maximum."
            )
        elif proposed_credits == preferred_max:
            recommendations.append(
                "The schedule is at the student's stated maximum, so avoid adding another course without reconsidering the overall workload."
            )

    # Stored combination information. Do not treat the manual rating as a validated probability.
    elevated = [p for p in matched_pairs if p.get("baseline_rating") is not None and p["baseline_rating"] >= 8]
    moderate = [p for p in matched_pairs if p.get("baseline_rating") is not None and 6 <= p["baseline_rating"] < 8]
    if elevated:
        pair_names = [f"{display_course_code(p['course_a'])} + {display_course_code(p['course_b'])}" for p in elevated]
        concerns.append("The proposed schedule contains an elevated baseline course-combination consideration: " + "; ".join(pair_names) + ".")
        recommendations.append(
            "Review the reasoning for the elevated course combination and consider discussing that pairing with an advisor before registration."
        )
    elif moderate:
        recommendations.append(
            "Pay particular attention to the stored course-combination notes when planning weekly study time."
        )

    # Recent experience with proposed courses can be useful context.
    repeated = sorted(set(recent_courses) & set(proposed_courses))
    if repeated:
        strengths.append(
            "The student has recent experience with " + ", ".join(display_course_code(c) for c in repeated) + "."
        )

    # Produce an overall action recommendation from the combined evidence.
    high_concern = (
        (difficulty >= 4 and satisfied in {"mostly", "no"})
        or (preferred_max is not None and proposed_credits > preferred_max)
        or (work_hours >= 20 and difficulty >= 4)
    )

    if high_concern:
        recommendations.insert(0,
            "Based on the student's recent experience and current circumstances, consider adjusting the proposed schedule or comparing an alternative schedule before registration."
        )
    elif difficulty <= 2 and satisfied == "yes" and (recent_credits is None or proposed_credits <= recent_credits) and not elevated:
        recommendations.insert(0,
            "Based on the student's reported experience, the proposed load appears reasonably consistent with a workload the student has previously managed successfully."
        )
    else:
        if matched_pairs:
            recommendations.insert(0,
                "The proposed schedule may be manageable, but the student should consider the stored course-combination information, recent workload experience, and current time commitments together."
            )
        else:
            recommendations.insert(0,
                "Based on the information currently available, the proposed schedule may be manageable. This assessment is based primarily on the student's recent workload experience and current circumstances because no validated course-combination relationships were found for this schedule."
            )

    # Include concise reasons so the recommendation is explainable.
    if concerns:
        recommendations.append("Factors increasing concern: " + " ".join(concerns))
    if strengths:
        recommendations.append("Factors supporting the schedule: " + " ".join(strengths))

    return recommendations

def main():
    print("=" * 68)
    print("SMARTADVISOR — PERSONALIZED COURSE COMPATIBILITY PROTOTYPE")
    print("=" * 68)
    print("Testing uses synthetic profiles and the supplied course-data files.\n")

    try:
        profiles = load_profiles()
        catalog = load_catalog()
        pair_data, data_warnings = load_course_combinations(catalog)
    except (FileNotFoundError, PermissionError) as exc:
        print(f"ERROR: {exc}")
        print("Place this script in the same folder as the three required Excel files.")
        return

    print(f"Loaded {len(profiles)} synthetic profiles.")
    print(f"Loaded {len(catalog)} unique courses from the catalog.")
    print(f"Loaded {len(pair_data)} course-combination records.\n")
    if data_warnings:
        print("Data note:")
        for warning in data_warnings:
            print(" - " + warning)
        print()

    student = select_profile(profiles)
    completed = extract_completed_codes(student["completed_courses_raw"], catalog)

    print("\nPROFILE FOUND")
    print(f"Name: {student['name']}")
    print(f"Synthetic ID: {student['student_id']}")
    print(f"Major: {student['major']}")
    print(f"Standing: {student['standing']}")
    print(f"Recognized completed courses: {len(completed)}")
    if student["workdays"] or student["work_schedule"]:
        print(f"Stored work schedule: {student['workdays'] or 'Not listed'} | {student['work_schedule'] or 'Time not listed'}")
    else:
        print("Stored work schedule: Not listed")

    print("\nCURRENT STUDENT CONTEXT")
    difficulty = prompt_int(
        "How difficult did your most recent semester feel?\n"
        "  1 = Very manageable\n  2 = Manageable\n  3 = Moderately difficult\n"
        "  4 = Difficult\n  5 = Very difficult\nChoice: ", 1, 5
    )

    while True:
        satisfaction = input(
            "Were you satisfied with your academic performance during that workload? "
            "(Yes/Mostly/No): "
        ).strip().lower()
        if satisfaction in {"yes", "mostly", "no"}:
            break
        print("Please enter Yes, Mostly, or No.")

    recent_courses = select_recent_courses(catalog, completed)
    calculated_recent_credits = sum(catalog[c]["credits"] for c in recent_courses) if recent_courses else None
    if calculated_recent_credits is not None:
        print(f"SmartAdvisor found {calculated_recent_credits:g} catalog credits in the recent-semester courses entered.")

    recent_credits = calculated_recent_credits
    if recent_credits is None:
        recent_credits = prompt_float(
            "About how many credits did you take in that recent semester? "
            "(press Enter if unknown): ", 0, 30, allow_blank=True, default=None
        )
    current_work_hours = prompt_float(
        "About how many hours per week do you expect to work during the proposed semester? ", 0, 80
    )
    preferred_max = prompt_float(
        "What is the maximum number of credits you would prefer to take? "
        "(press Enter if you do not have a preference): ", 1, 30, allow_blank=True, default=None
    )

    selected = select_courses(catalog)
    proposed_credits = sum(catalog[c]["credits"] for c in selected)

    print("\n" + "=" * 68)
    print("PROPOSED SEMESTER")
    print("=" * 68)
    for code in selected:
        info = catalog[code]
        print(f"- {display_course_code(code)} — {info['title']} ({info['credits']:g} credits)")
    print(f"Total catalog credits: {proposed_credits:g}")

    already_completed = sorted(set(selected) & set(completed))
    if already_completed:
        print("\nPROFILE CHECK")
        print("The following proposed course(s) also appear in this synthetic student's completed-course history:")
        print("  " + ", ".join(display_course_code(c) for c in already_completed))
        print("Verify whether the student intends to repeat these courses. SmartAdvisor will continue")
        print("the compatibility test, but this should be reviewed before registration.")

    matched = []
    unmatched_count = 0
    print("\nCOURSE-COMBINATION REVIEW")
    if len(selected) < 2:
        print("Only one course was selected, so there are no course pairs to compare.")
    else:
        for a, b in combinations(selected, 2):
            pair = pair_data.get(frozenset((a, b)))
            if not pair:
                unmatched_count += 1
                continue

            matched.append(pair)
            result, factors = personalized_pair_guidance(
                pair, difficulty, satisfaction, current_work_hours,
                proposed_credits, preferred_max, recent_courses, selected
            )
            print(f"\n{display_course_code(a)} + {display_course_code(b)}")
            print(f"  Baseline: {baseline_label(pair['baseline_rating'])}")
            print(f"  Course-data reasoning: {pair['reasoning']}")
            print(f"  Personalized guidance: {result}")
            for factor in factors:
                print(f"   • {factor}")

        if not matched:
            print("No validated course-combination relationships were found for this proposed schedule")
            print("in the current comparison dataset. The semester assessment will therefore rely")
            print("primarily on the student's recent workload experience and current circumstances.")
        elif unmatched_count:
            print(f"\nData note: {unmatched_count} additional course pair(s) do not yet have a validated")
            print("comparison in the current dataset and were not used to create a pair-specific assessment.")

    overall = overall_semester_guidance(
        proposed_credits, preferred_max, difficulty, satisfaction,
        current_work_hours, len(matched)
    )

    print("\n" + "=" * 68)
    print("PERSONALIZED SEMESTER SUMMARY")
    print("=" * 68)
    print(f"Assessment: {overall}")
    print(f"Matched stored course combinations: {len(matched)}")
    if recent_courses:
        print("Recent-semester courses used for comparison: " +
              ", ".join(display_course_code(c) for c in recent_courses))
    if recent_credits is not None:
        difference = proposed_credits - recent_credits
        if difference > 0:
            print(f"The proposed {proposed_credits:g}-credit semester is {difference:g} credits higher than the student's recent {recent_credits:g}-credit semester.")
        elif difference < 0:
            print(f"The proposed {proposed_credits:g}-credit semester is {abs(difference):g} credits lower than the student's recent {recent_credits:g}-credit semester.")
        else:
            print(f"The proposed {proposed_credits:g}-credit semester has the same credit load as the student's recent semester.")
    if preferred_max is not None:
        print(f"Student's preferred maximum: {preferred_max:g} credits")

    print("\nPERSONALIZED RECOMMENDATION")
    recommendation_lines = build_personalized_recommendation(
        proposed_credits=proposed_credits,
        recent_credits=recent_credits,
        preferred_max=preferred_max,
        difficulty=difficulty,
        satisfied=satisfaction,
        work_hours=current_work_hours,
        matched_pairs=matched,
        recent_courses=recent_courses,
        proposed_courses=selected,
    )
    for line in recommendation_lines:
        print("- " + line)

    print("\nImportant: This prototype provides preliminary planning guidance only.")
    print("It does not determine whether a student should or should not take a course,")
    print("and it does not replace official FAU advising, prerequisites, or registration rules.")


if __name__ == "__main__":
    main()
